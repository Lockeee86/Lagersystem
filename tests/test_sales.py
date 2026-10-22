import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import sales, storage
from app.models import Certification, Customer, Delivery, DeliveryItem, Movement, SeedClass, Setting
from app.storage import StorageError


@pytest.fixture
def customer(db, seed):
    c = Customer(name="Saatzucht Nord", number="4711", address="Hofstraße 2\n12345 Nordheim")
    db.session.add(c)
    db.session.commit()
    return c


def filled_boxes(db, seed, n, fill=100, stack=(1, 1)):
    boxes = storage.create_boxes(n, seed["bt"].id, seed["admin"])
    for box in boxes:
        storage.update_contents(box, seed["admin"], seed["belana"].id, date(2026, 9, 23), seed["plot"].id,
                                "35/55", fill, None, usage="pflanzgut")
        storage.store_box(box, seed["wh"].id, *stack, seed["admin"])
    db.session.commit()
    return boxes


def test_numbers_per_year(db, seed, customer):
    a = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    b = sales.create_delivery(customer, date(2026, 10, 2), seed["admin"])
    c = sales.create_delivery(customer, date(2027, 1, 5), seed["admin"])
    assert (a.number, b.number, c.number) == ("LS-2026-0001", "LS-2026-0002", "LS-2027-0001")


def test_load_box_takes_it_from_stack_and_empties_it(db, seed, customer):
    a, b, c = filled_boxes(db, seed, 3)
    delivery = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    item, shifted = sales.add_box(delivery, a, seed["worker"])
    db.session.commit()
    assert shifted == [b, c] and (b.level, c.level) == (1, 2)
    assert not a.is_stored and a.fill == 0
    assert item.fill == 100 and item.weight_kg == 400 and item.from_location.endswith("A-01-1")
    assert item.variety_id == seed["belana"].id and item.usage == "pflanzgut" and item.grade == "35/55"
    m = a.movements.first()
    assert m.action == "verladen" and m.detail == delivery.number and m.fill_before == 100
    with pytest.raises(StorageError):   # leer → nicht noch einmal verladbar
        sales.add_box(delivery, a, seed["worker"])


def test_remove_item_restores_fill(db, seed, customer):
    (a,) = filled_boxes(db, seed, 1, fill=75)
    delivery = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    item, _ = sales.add_box(delivery, a, seed["worker"])
    db.session.commit()
    sales.remove_item(item, seed["worker"])
    db.session.commit()
    assert a.fill == 75 and not a.is_stored and DeliveryItem.query.count() == 0
    assert a.movements.first().action == "zurueckgenommen"


def test_summary_groups_and_certification(db, seed, customer):
    pb = SeedClass(name="PB")
    db.session.add(pb)
    db.session.commit()
    boxes = filled_boxes(db, seed, 4)
    boxes[3].fill = 50
    for box in boxes:
        box.seed_class_id = pb.id
    db.session.add(Certification(variety_id=seed["belana"].id, seed_class_id=pb.id, year=2026, number="DE-123"))
    delivery = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    for box in boxes:
        sales.add_box(delivery, box, seed["worker"])
    sales.add_quantity(delivery, seed["admin"], seed["belana"].id, "speise", None, None, "25-kg-Sack", 40, 25)
    db.session.commit()
    lines = sales.summarize(delivery)
    assert len(lines) == 2
    box_line, bag_line = lines
    assert box_line["boxes"] == 4 and box_line["box_units"] == 3.5 and box_line["weight_kg"] == 1400
    assert box_line["code_ranges"] == "K-0001 – K-0004"
    assert box_line["certification"].number == "DE-123"
    assert bag_line["quantity"] == 40 and bag_line["weight_kg"] == 1000 and not bag_line["is_box"]
    assert delivery.estimated_kg == 2400


def test_code_ranges():
    assert sales.code_ranges(["K-0003", "K-0001", "K-0002", "K-0007", "K-0009", "K-0008"]) == \
        "K-0001 – K-0003, K-0007 – K-0009"
    assert sales.code_ranges([]) == ""


def test_closed_delivery_is_locked(db, seed, customer):
    (a, b) = filled_boxes(db, seed, 2)
    delivery = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    with pytest.raises(StorageError):
        sales.close_delivery(delivery)          # ohne Positionen
    sales.add_box(delivery, a, seed["worker"])
    sales.close_delivery(delivery)
    with pytest.raises(StorageError):
        sales.add_box(delivery, b, seed["worker"])
    with pytest.raises(StorageError):
        sales.set_weights(delivery, 1000, 500)
    with pytest.raises(StorageError):
        sales.delete_delivery(delivery)


def test_weights(db, seed, customer):
    delivery = sales.create_delivery(customer, date(2026, 10, 1), seed["admin"])
    with pytest.raises(StorageError):
        sales.set_weights(delivery, 1000, 3000)
    sales.set_weights(delivery, 18400, 9200)
    assert delivery.net_kg == 9200


# --- Seiten -----------------------------------------------------------------

def test_flow_via_pages(as_admin, as_worker, as_viewer, db, seed, customer):
    a, b = filled_boxes(db, seed, 2)
    Setting.put("farm_name", "Hof Sonnenfeld")
    Setting.put("farm_passport_reg", "DE-NI1-12345")
    db.session.commit()

    assert as_worker.post("/lieferscheine/neu", data={"customer_id": customer.id}).status_code == 403
    resp = as_admin.post("/lieferscheine/neu", data={"customer_id": customer.id, "date": "2026-10-20",
                                                     "vehicle": "AB-C 123"})
    delivery = Delivery.query.one()
    assert resp.headers["Location"].endswith(f"/lieferschein/{delivery.id}")

    # Mitarbeiter scannt (JSON wie der Dauer-Scanner)
    json = {"Accept": "application/json"}
    r = as_worker.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "1"}, headers=json).get_json()
    assert r["ok"] and r["code"] == "K-0001" and r["count_text"] == "1 Kiste"
    r = as_worker.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "K-0001"}, headers=json).get_json()
    assert not r["ok"] and "LS-2026-0001" in r["message"]
    r = as_worker.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "999"}, headers=json).get_json()
    assert not r["ok"]
    as_worker.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "K-0002"})   # Formular
    assert delivery.box_count == 2
    assert as_viewer.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "K-0002"}).status_code == 403

    for url in (f"/lieferschein/{delivery.id}", f"/lieferschein/{delivery.id}/verladen", "/lieferscheine", "/"):
        assert as_worker.get(url).status_code == 200, url
    assert "LS-2026-0001" in as_viewer.get(f"/kiste/{a.code}").get_data(as_text=True)

    as_admin.post(f"/lieferschein/{delivery.id}/waage", data={"gross_kg": "12000", "tare_kg": "11250"})
    as_admin.post(f"/lieferschein/{delivery.id}/menge", data={"variety_id": seed["belana"].id, "usage": "speise",
                                                               "quantity": "10", "packaging": "25-kg-Sack",
                                                               "unit_kg": "25"})
    assert delivery.net_kg == 750 and delivery.estimated_kg == 1050

    pdf = as_viewer.get(f"/lieferschein/{delivery.id}.pdf")
    assert pdf.status_code == 200 and pdf.data[:4] == b"%PDF"

    as_admin.post(f"/lieferschein/{delivery.id}/abschliessen")
    assert not delivery.is_open
    r = as_worker.post(f"/lieferschein/{delivery.id}/kiste", data={"code": "K-0001"}, headers=json).get_json()
    assert not r["ok"]
    assert as_worker.get(f"/lieferschein/{delivery.id}/verladen").status_code == 302

    # Export enthält die Positionen
    wb = load_workbook(io.BytesIO(as_admin.get("/admin/daten/export.xlsx?jahr=2026").data))
    rows = list(wb["Lieferscheine"].iter_rows(values_only=True))
    assert len(rows) == 4 and rows[1][0] == "LS-2026-0001"

    # Kisten mit Lieferschein lassen sich nicht löschen, Kunden nur deaktivieren
    as_admin.post(f"/kiste/{a.code}/loeschen")
    assert db.session.get(type(a), a.id) is not None
    as_admin.post(f"/kunden/{customer.id}/loeschen")
    assert customer.active is False


def test_customer_and_certification_admin(as_admin, as_worker, db, seed):
    assert as_worker.get("/kunden").status_code == 403
    as_admin.post("/kunden", data={"name": "Händler Süd", "address": "Weg 1"})
    as_admin.post("/kunden", data={"name": "händler süd"})
    assert Customer.query.count() == 1
    pb = SeedClass(name="PB")
    db.session.add(pb)
    db.session.commit()
    data = {"variety_id": seed["belana"].id, "seed_class_id": pb.id, "year": "2026", "number": "DE-1"}
    as_admin.post("/anerkennungen", data=data)
    as_admin.post("/anerkennungen", data=data | {"number": "DE-2"})   # doppelt → abgelehnt
    assert [c.number for c in Certification.query] == ["DE-1"]
    for url in ("/kunden", "/anerkennungen"):
        assert as_admin.get(url).status_code == 200
