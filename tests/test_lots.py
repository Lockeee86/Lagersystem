import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app import storage, transfer
from app.models import Lot, Movement, SeedClass, Variety, format_quantity
from app.storage import StorageError

D = Decimal


@pytest.fixture
def be(db):
    seed_class = SeedClass(name="BE", position=1)
    db.session.add(seed_class)
    db.session.commit()
    return seed_class


def new_lot(seed, row=6, qty="52.3", **kw):
    return storage.create_lot(seed["admin"], seed["wh"].id, row, seed["belana"].id, D(qty), seed["bt"].id, **kw)


# --- Lagerlogik ---------------------------------------------------------------

def test_parse_and_format_quantity():
    assert storage.parse_quantity("41,75") == D("41.75")
    assert storage.parse_quantity("40x") == D("40.00")
    assert storage.parse_quantity(57.1) == D("57.10")
    assert storage.parse_quantity("abc") is None
    assert [format_quantity(v) for v in (D("41.75"), D("40.00"), D("57.10"))] == ["41,75", "40", "57,1"]


def test_create_take_add_correct(db, seed, be):
    lot = new_lot(seed, qty="41.75", seed_class_id=be.id, usage="pflanzgut", harvest_date=date(2026, 9, 22))
    db.session.commit()
    assert (lot.title, lot.quantity_label, lot.position, lot.weight_kg) == ("Belana Saat BE", "41,75", 1, 16700)

    storage.take_from_lot(lot, D("5"), seed["worker"], note="Hofladen")
    storage.add_to_lot(lot, D("0.25"), seed["worker"])
    assert lot.quantity == D("37.00")
    with pytest.raises(StorageError, match="nur 37"):
        storage.take_from_lot(lot, D("40"), seed["worker"])
    assert storage.correct_lot(lot, D("36.5"), seed["worker"])
    storage.take_from_lot(lot, D("36.5"), seed["worker"])
    db.session.commit()
    assert lot.closed and lot.quantity == 0

    moves = Movement.query.filter_by(lot_id=lot.id).order_by(Movement.id).all()
    assert [m.action for m in moves] == ["partie_angelegt", "entnommen", "zugang", "korrektur", "entnommen"]
    assert [m.quantity_after for m in moves] == [D("41.75"), D("36.75"), D("37"), D("36.5"), D("0")]
    assert all(m.variety_id == seed["belana"].id and m.seed_class_id == be.id for m in moves)
    assert moves[1].note == "Hofladen" and moves[1].quantity_delta_label == "−5"


def test_positions_and_shift(db, seed):
    a, b, c = new_lot(seed, qty="40"), new_lot(seed, qty="9"), new_lot(seed, qty="20")
    db.session.commit()
    assert [x.position for x in (a, b, c)] == [1, 2, 3]
    assert storage.shift_lot(c, -1)
    assert [x.id for x in storage.row_lots(seed["wh"].id, 6)] == [a.id, c.id, b.id]
    assert not storage.shift_lot(a, -1)
    assert storage.row_quantity(seed["wh"].id, 6) == D("69")


def test_move_whole_and_split(db, seed):
    lot = new_lot(seed, row=6, qty="52.3")
    db.session.commit()
    part = storage.move_lot(lot, seed["wh"].id, 2, seed["worker"], quantity=D("10"))
    db.session.commit()
    assert (lot.row, lot.quantity) == (6, D("42.30"))
    assert (part.row, part.quantity, part.variety_id) == (2, D("10.00"), lot.variety_id)
    assert Movement.query.filter_by(lot_id=lot.id, action="geteilt").one().detail == str(part.id)

    storage.move_lot(lot, seed["wh"].id, 3, seed["worker"])
    db.session.commit()
    assert lot.row == 3
    with pytest.raises(StorageError, match="bereits"):
        storage.move_lot(lot, seed["wh"].id, 3, seed["worker"])
    with pytest.raises(StorageError):
        storage.move_lot(lot, seed["wh"].id, 11, seed["worker"])


def test_update_lot_logs_changes(db, seed, be):
    lot = new_lot(seed)
    db.session.commit()
    assert storage.update_lot(lot, seed["worker"], seed["belana"].id, "speise", be.id, "35/50", None,
                              date(2026, 10, 19), None, seed["bt"].id)
    db.session.commit()
    m = Movement.query.filter_by(lot_id=lot.id, action="partie_geaendert").one()
    assert m.detail == "usage,seed_class,grade,harvest_date"
    assert m.detail_text == "Geändert: Verwendung, Stufe, Sortierung, Erntedatum"
    assert lot.title == "Belana SP BE 35/50"


# --- Seiten -----------------------------------------------------------------

def test_tafel_view_and_lot_pages(as_worker, db, seed, be):
    seed["wh"].numeric_rows = True
    lot = new_lot(seed, row=9, qty="40", seed_class_id=be.id, harvest_date=date(2026, 9, 23))
    new_lot(seed, row=9, qty="34")
    db.session.commit()

    html = as_worker.get("/lager?ansicht=tafel").get_data(as_text=True)
    assert 'class="tafel"' in html and "40×" in html and "Belana" in html and "Saat BE" in html and "23.9." in html
    assert ">9<" in html and "74</span>" in html and "1 frei" in html
    assert html.index(">10<") < html.index(">9<") < html.index(">1<")   # höchste Reihe oben

    html = as_worker.get(f"/partie/{lot.id}").get_data(as_text=True)
    assert "Entnehmen" in html and "Reihe 9" in html

    resp = as_worker.post(f"/partie/{lot.id}/entnehmen", data={"quantity": "2,5", "note": "Kunde"}, follow_redirects=True)
    assert "2,5 Kisten entnommen. Bestand: 37,5." in resp.get_data(as_text=True)
    as_worker.post(f"/partie/{lot.id}/umsetzen", data={"warehouse_id": seed["wh"].id, "row": 4, "quantity": "7,5"})
    assert lot.quantity == D("30.00")
    assert Lot.query.filter_by(row=4).one().quantity == D("7.50")
    as_worker.post(f"/partie/{lot.id}/korrektur", data={"quantity": "29"})
    assert lot.quantity == D("29.00")
    resp = as_worker.post(f"/partie/{lot.id}/entnehmen", data={"quantity": "x"}, follow_redirects=True)
    assert "gültige Menge" in resp.get_data(as_text=True)

    # Verlauf zeigt Partien
    html = as_worker.get("/verlauf?q=Belana").get_data(as_text=True)
    assert "Entnommen" in html and "−2,5 Kisten" in html


def test_create_and_edit_lot_via_form(as_worker, db, seed, be):
    resp = as_worker.post("/partie/neu", data={
        "warehouse_id": seed["wh"].id, "row": 5, "quantity": "42", "variety_id": seed["belana"].id,
        "usage": "pflanzgut", "seed_class_id": be.id, "harvest_date": "2026-09-10", "box_type_id": seed["bt"].id,
    })
    assert resp.status_code == 302
    lot = Lot.query.one()
    assert (lot.row, lot.quantity, lot.seed_class_id) == (5, D("42.00"), be.id)
    as_worker.post(f"/partie/{lot.id}/bearbeiten", data={
        "variety_id": seed["belana"].id, "usage": "speise", "box_type_id": seed["bt"].id, "grade": "35/50"})
    assert (lot.usage, lot.grade, lot.seed_class_id, lot.harvest_date) == ("speise", "35/50", None, None)
    # Ohne Sorte → Fehler, nichts angelegt
    html = as_worker.post("/partie/neu", data={"warehouse_id": seed["wh"].id, "row": 5, "quantity": "3"}).get_data(as_text=True)
    assert "Bitte eine Sorte wählen." in html and Lot.query.count() == 1


def test_viewer_cannot_change_lots(as_viewer, db, seed):
    lot = new_lot(seed)
    db.session.commit()
    assert as_viewer.get(f"/partie/{lot.id}").status_code == 200
    assert as_viewer.post(f"/partie/{lot.id}/entnehmen", data={"quantity": "1"}).status_code == 403
    assert as_viewer.get("/partie/neu").status_code == 403


def test_dashboard_counts_lots(as_admin, db, seed):
    new_lot(seed, qty="41.75", usage="pflanzgut")
    new_lot(seed, row=1, qty="24", usage="speise")
    db.session.commit()
    html = as_admin.get("/").get_data(as_text=True)
    assert "65,75 befüllte Kisten" in html and "Saat" in html and "Speise" in html
    assert "26,3 t" in html   # 65,75 × 400 kg


def test_seed_class_admin(as_admin, db, seed):
    as_admin.post("/admin/stufen", data={"name": "PB", "description": "Vorstufe", "position": 1})
    assert SeedClass.query.filter_by(name="PB").one().description == "Vorstufe"
    pb = SeedClass.query.filter_by(name="PB").one()
    new_lot(seed, seed_class_id=pb.id)
    db.session.commit()
    as_admin.post(f"/admin/stufen/{pb.id}/loeschen")
    assert pb.active is False


def test_numeric_rows(as_admin, db, seed):
    as_admin.post("/admin/lager", data={"id": seed["wh"].id, "name": "Halle 1", "rows": 9, "slots": 15,
                                        "max_levels": 5, "numeric_rows": "1"})
    assert seed["wh"].numeric_rows and seed["wh"].row_label(3) == "3" and seed["wh"].rows == 9


# --- Tafel-Import / Export ----------------------------------------------------

def tafel_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Partien"
    ws.append(["Reihe", "Menge", "Sorte", "Verwendung", "Stufe", "Kaliber", "Datum", "Schlag", "Lager", "Notiz"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_lot_template_round_trip(as_admin, seed):
    data = as_admin.get("/admin/daten/tafel-vorlage.xlsx").data
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Anleitung", "Partien", "Listen"]
    assert transfer.read_file("t.xlsx", data) == ("lots", [])


def test_lot_import_from_board(as_admin, db, seed):
    data = tafel_xlsx([
        ("9", "40", "Rumba", "Saat", "BE", None, "23.9", None, None, None),
        (9, "9", "Rumba", "Saat", "PB", None, "23.9.2026", None, None, None),
        ("8", "41,75", "Rumba", "Saat", "Z", None, "22.09.2026", None, None, None),
        ("1", "24", "Laura", "SP", None, "35/50", "19.10", "Acker Nord", None, "vorne"),
    ])
    html = as_admin.post("/admin/daten/import", data={"file": (io.BytesIO(data), "tafel.xlsx")},
                         content_type="multipart/form-data").get_data(as_text=True)
    assert "Sorte „Rumba“ gibt es nicht." in html and "Import übernehmen" not in html

    html = as_admin.post("/admin/daten/import", data={"file": (io.BytesIO(data), "tafel.xlsx"), "create_missing": "1"},
                         content_type="multipart/form-data").get_data(as_text=True)
    assert "Import übernehmen" in html and "Partie 41,75× Rumba Saat Z angelegt" in html
    assert Lot.query.count() == 0   # nur Vorschau
    as_admin.post("/admin/daten/import/uebernehmen")
    lots = Lot.query.order_by(Lot.row.desc(), Lot.position).all()
    assert [(lot.row, lot.position, lot.quantity_label, lot.title) for lot in lots] == [
        (9, 1, "40", "Rumba Saat BE"), (9, 2, "9", "Rumba Saat PB"), (8, 1, "41,75", "Rumba Saat Z"),
        (1, 1, "24", "Laura SP 35/50")]
    assert lots[0].harvest_date.month == 9 and lots[0].harvest_date.day == 23
    assert lots[3].plot.name == "Acker Nord" and lots[3].note == "vorne"
    assert Variety.query.filter_by(name="Rumba").one().color


def test_lot_import_errors(as_admin, db, seed):
    data = tafel_xlsx([("12", "5", "Belana", None, None, None, None, None, None, None),
                       ("3", "viele", "Belana", None, None, None, None, None, None, None),
                       ("3", "5", "Belana", "Gemüse", None, None, "32.13", None, None, None)])
    html = as_admin.post("/admin/daten/import", data={"file": (io.BytesIO(data), "t.xlsx")},
                         content_type="multipart/form-data").get_data(as_text=True)
    assert "Reihe „12“ gibt es" in html and "Menge „viele“" in html and "Verwendung „Gemüse“" in html


def test_season_date_guess():
    assert transfer._season_date(23, 9, today=date(2026, 10, 22)) == date(2026, 9, 23)
    assert transfer._season_date(23, 9, today=date(2027, 3, 1)) == date(2026, 9, 23)   # Winter: Vorjahr
    with pytest.raises(transfer.RowError):
        transfer._parse_date(20.1)   # Excel-Zahl aus „20.10“ ist mehrdeutig


def test_export_with_lots(as_admin, db, seed):
    lot = new_lot(seed, qty="40", usage="pflanzgut", harvest_date=date(2026, 9, 23), plot_id=seed["plot"].id)
    storage.take_from_lot(lot, D("10"), seed["admin"])
    db.session.commit()
    groups = transfer.season_summary(2026)
    g = groups[(2026, seed["belana"], "pflanzgut", None)]
    assert (g["boxes"], g["harvest_kg"], g["stock_kg"], g["stored"]) == (40, 16000, 12000, 30)
    wb = load_workbook(io.BytesIO(as_admin.get("/admin/daten/export.xlsx?jahr=2026").data))
    assert wb.sheetnames == ["Auswertung", "Partien", "Kisten", "Bewegungen"]
    assert list(wb["Partien"].iter_rows(values_only=True))[1][:7] == ("Kartoffellager", "F", 1, 30, "Belana", "Saat", None)
    values = [tuple(r[:8]) for r in wb["Auswertung"].iter_rows(values_only=True)]
    assert (2026, "Belana", "Saat", None, 40, 16000, 12000, 30) in values
    assert (2026, "Belana", "Saat", "Acker Nord", 40, 16000, 12000, 30) in values
    moves = list(wb["Bewegungen"].iter_rows(values_only=True))
    assert [(m[2], m[7], m[8]) for m in moves[1:]] == [("Partie angelegt", 40, 40), ("Entnommen", -10, 30)]
