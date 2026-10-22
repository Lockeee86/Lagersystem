import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app import storage, transfer
from app.models import Box, Movement, Plot, Variety


def xlsx(rows, headers=("Kiste", "Sorte", "Schlag", "Erntedatum", "Füllstand", "Reihe", "Platz", "Auslagern", "Notiz")):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kisten"
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, data, name="zettel.xlsx", create_missing=False):
    form = {"file": (io.BytesIO(data), name)}
    if create_missing:
        form["create_missing"] = "1"
    return client.post("/admin/daten/import", data=form, content_type="multipart/form-data")


def make_boxes(db, seed, n):
    boxes = storage.create_boxes(n, seed["bt"].id, seed["admin"])
    db.session.commit()
    return boxes


# --- Vorlage ----------------------------------------------------------------

def test_template_has_lists_and_reads_back_empty(as_admin, seed):
    resp = as_admin.get("/admin/daten/vorlage.xlsx")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Anleitung", "Kisten", "Listen"]
    ws = wb["Kisten"]
    assert [c.value for c in ws[1]][:3] == ["Kiste", "Sorte", "Schlag"]
    assert ws.data_validations.dataValidation  # Auswahllisten vorhanden
    assert wb["Listen"]["A2"].value == "Belana" and wb["Listen"]["B2"].value == "Acker Nord"
    # Leere Vorlage → keine Zeilen
    assert transfer.read_file("v.xlsx", resp.data) == ("boxes", [])


def test_template_in_polish_can_be_read(as_admin, db, seed):
    make_boxes(db, seed, 1)
    as_admin.post("/sprache", data={"lang": "pl"})
    data = as_admin.get("/admin/daten/vorlage.xlsx").data
    wb = load_workbook(io.BytesIO(data))
    assert "Skrzynie" in wb.sheetnames
    ws = wb["Skrzynie"]
    ws["A2"] = "1"
    ws["E2"] = "pełna"
    buf = io.BytesIO()
    wb.save(buf)
    kind, rows = transfer.read_file("pl.xlsx", buf.getvalue())
    assert kind == "boxes"
    assert rows == [(2, {"code": "1", "variety": None, "plot": None, "harvest_date": None, "fill": "pełna",
                         "grade": None, "warehouse": None, "row": None, "slot": None, "remove": None, "note": None})]


# --- Import -----------------------------------------------------------------

def test_import_preview_then_apply(as_admin, db, seed):
    make_boxes(db, seed, 3)
    data = xlsx([
        ("K-0001", "Belana", "Acker Nord", date(2026, 10, 5), "voll", "C", 7, None, "erste"),
        ("2", "belana", None, "06.10.2026", "¾", "C", 7, None, None),
        (3, None, None, None, None, None, None, None, None),
    ])
    resp = upload(as_admin, data)
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200 and "Import übernehmen" in html
    assert "eingelagert auf C-07-1" in html and "eingelagert auf C-07-2" in html and "keine Änderung" in html
    # Vorschau speichert nichts
    assert Box.query.filter(Box.warehouse_id.isnot(None)).count() == 0

    resp = as_admin.post("/admin/daten/import/uebernehmen", follow_redirects=True)
    assert "3 Zeilen importiert." in resp.get_data(as_text=True)
    b1, b2, b3 = Box.query.order_by(Box.code).all()
    assert (b1.location, b1.variety.name, b1.plot.name, b1.fill, b1.note) == ("C-07-1", "Belana", "Acker Nord", 100, "erste")
    assert (b2.location, b2.fill, b2.harvest_date) == ("C-07-2", 75, date(2026, 10, 6))
    assert not b3.is_stored
    # Datei ist nach dem Übernehmen weg → zweites Übernehmen geht nicht
    resp = as_admin.post("/admin/daten/import/uebernehmen", follow_redirects=True)
    assert "abgelaufen" in resp.get_data(as_text=True)


def test_import_with_errors_saves_nothing(as_admin, db, seed):
    make_boxes(db, seed, 2)
    data = xlsx([
        ("1", "Belana", None, None, None, "A", 1, None, None),
        ("2", "Gibtsnicht", None, None, "sehr voll", "Z", 1, None, None),
        ("99", None, None, None, None, None, None, None, None),
    ])
    html = upload(as_admin, data).get_data(as_text=True)
    assert "Import übernehmen" not in html
    assert "Sorte „Gibtsnicht“ gibt es nicht." in html
    assert "K-0099 gibt es nicht." in html
    assert Box.query.filter(Box.warehouse_id.isnot(None)).count() == 0
    # Auch ein direktes Übernehmen speichert nichts
    as_admin.post("/admin/daten/import/uebernehmen")
    assert Box.query.filter(Box.warehouse_id.isnot(None)).count() == 0


def test_import_creates_missing_boxes_and_removes(as_admin, db, seed):
    (box,) = make_boxes(db, seed, 1)
    storage.update_contents(box, seed["admin"], seed["belana"].id, None, None, None, 100, None)
    storage.store_box(box, seed["wh"].id, 1, 1, seed["admin"])
    db.session.commit()
    data = xlsx([
        ("K-0001", None, None, None, "leer", None, None, "x", None),
        ("K-0500", "Belana", None, None, None, "B", 2, None, None),
    ])
    upload(as_admin, data, create_missing=True)
    as_admin.post("/admin/daten/import/uebernehmen")
    assert not box.is_stored and box.fill == 0
    new = Box.query.filter_by(code="K-0500").one()
    assert new.location == "B-02-1" and new.fill == 100   # leere Kiste mit Sorte → voll


def test_import_csv_with_semicolons(as_admin, db, seed):
    make_boxes(db, seed, 1)
    csv_data = "Kistennummer;Sorte;Füllstand;Reihe;Platz\nK-0001;Belana;50%;D;3\n".encode("cp1252")
    upload(as_admin, csv_data, name="zettel.csv")
    as_admin.post("/admin/daten/import/uebernehmen")
    box = Box.query.one()
    assert box.location == "D-03-1" and box.fill == 50


def test_import_rejects_wrong_files(as_admin, seed):
    resp = upload(as_admin, b"hallo", name="bild.png")
    assert resp.status_code == 302
    assert "Bitte eine Excel-Datei" in as_admin.get("/admin/daten").get_data(as_text=True)
    resp = upload(as_admin, xlsx([("1", "x")], headers=("Foo", "Bar")))
    assert resp.status_code == 302
    assert "Keine passende Kopfzeile" in as_admin.get("/admin/daten").get_data(as_text=True)


def test_parse_fill_variants():
    assert transfer._parse_fill(0.75) == 75   # Excel-Prozentzelle
    assert transfer._parse_fill(1) == 100
    assert transfer._parse_fill("1/2") == 50
    assert transfer._parse_fill("25 %") == 25
    assert transfer._parse_fill("pusta") == 0


def test_import_only_for_admins(as_worker, seed):
    assert as_worker.get("/admin/daten").status_code == 403
    assert as_worker.get("/admin/daten/export.xlsx").status_code == 403


# --- Export / Kartoffeljahr ---------------------------------------------------

def test_snapshot_keeps_old_season_after_refill(as_admin, db, seed):
    (box,) = make_boxes(db, seed, 1)
    linda = Variety(name="Linda", color="#c8553d")
    hof = Plot(name="Hofwiese")
    db.session.add_all([linda, hof])
    db.session.commit()

    # Saison 2026: Belana vom Acker Nord, voll eingelagert, später halb
    storage.update_contents(box, seed["admin"], seed["belana"].id, date(2026, 10, 1), seed["plot"].id, None, 100, None)
    storage.store_box(box, seed["wh"].id, 1, 1, seed["admin"])
    storage.set_fill(box, 50, seed["admin"])
    storage.remove_box(box, seed["admin"], empty=True)
    # Saison 2027: dieselbe Kiste neu befüllt mit Linda
    storage.update_contents(box, seed["admin"], linda.id, date(2027, 9, 20), hof.id, None, 75, None)
    db.session.commit()

    assert transfer.season_years() == [2027, 2026]
    groups = transfer.season_summary(2026)
    g = groups[(2026, seed["belana"], "speise", None)]
    assert (g["boxes"], g["harvest_kg"], g["stock_kg"]) == (1, 400, 0)
    assert groups[(2026, seed["belana"], "speise", seed["plot"])]["harvest_kg"] == 400
    g = transfer.season_summary(2027)[(2027, linda, "speise", None)]
    assert (g["boxes"], g["harvest_kg"], g["stock_kg"]) == (1, 300, 300)

    resp = as_admin.get("/admin/daten/export.xlsx?jahr=2026")
    assert resp.status_code == 200
    assert "lager_2026_" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Auswertung", "Partien", "Kisten", "Bewegungen"]
    values = [tuple(row[:8]) for row in wb["Auswertung"].iter_rows(values_only=True)]
    assert (2026, "Belana", "Speise", None, 1, 400, 0, 0) in values
    assert (2026, "Belana", "Speise", "Acker Nord", 1, 400, 0, 0) in values
    moves = list(wb["Bewegungen"].iter_rows(values_only=True))
    assert len(moves) == 1 + 4   # Kopfzeile + 4 Bewegungen der Saison 2026
    assert all(m[9] == "Belana" for m in moves[1:])
    # Kistenblatt 2026 ist leer (Kiste gehört jetzt zu 2027)
    assert wb["Kisten"].max_row == 1

    wb = load_workbook(io.BytesIO(as_admin.get("/admin/daten/export.xlsx").data))
    assert wb["Kisten"].max_row == 2 and wb["Bewegungen"].max_row == 1 + Movement.query.count()


def test_used_variety_and_plot_only_deactivated(as_admin, db, seed):
    (box,) = make_boxes(db, seed, 1)
    storage.update_contents(box, seed["admin"], seed["belana"].id, None, seed["plot"].id, None, 100, None)
    storage.update_contents(box, seed["admin"], None, None, None, None, 0, None)
    db.session.commit()
    # Kiste nutzt Sorte/Schlag nicht mehr, der Verlauf aber schon
    as_admin.post(f"/admin/sorten/{seed['belana'].id}/loeschen")
    as_admin.post(f"/admin/schlaege/{seed['plot'].id}/loeschen")
    assert seed["belana"].active is False and seed["plot"].active is False


def test_data_page_renders(as_admin, db, seed):
    (box,) = make_boxes(db, seed, 1)
    storage.update_contents(box, seed["admin"], seed["belana"].id, date(2026, 10, 1), None, None, 100, None)
    db.session.commit()
    html = as_admin.get("/admin/daten").get_data(as_text=True)
    assert '<option value="2026" selected>2026</option>' in html
