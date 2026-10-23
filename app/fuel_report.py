"""Tankbuch-Bericht für einen Zeitraum: als Excel-Datei und als PDF (z. B. für Agrardiesel-Erstattung, Steuerberater).

Wie der Lieferschein bewusst immer auf Deutsch – es ist ein Dokument für Dritte.
"""
import io
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import fuel
from .models import FuelEntry

KIND_TEXT = {"tanken": "Getankt", "lieferung": "Lieferung", "peilung": "Peilung"}


# --- Auswertung -----------------------------------------------------------------

def build(start, end, machine=None):
    """Alles für Excel und PDF: Einträge, Bestände, Summen, Werte je Maschine."""
    query = FuelEntry.query.filter(FuelEntry.date >= start, FuelEntry.date <= end)
    if machine is not None:                 # nur die Tankungen dieser Maschine (ohne Lieferungen/Peilungen)
        query = query.filter(FuelEntry.machine_id == machine.id)
    entries = query.order_by(FuelEntry.date, FuelEntry.id).all()
    refuels = [e for e in entries if e.kind == "tanken"]
    deliveries = [e for e in entries if e.kind == "lieferung"]

    # Preis: gewichteter Durchschnitt der Lieferungen im Zeitraum, sonst letzter Preis davor
    priced = [e for e in deliveries if e.price_per_liter]
    if priced:
        price = sum(e.liters * e.price_per_liter for e in priced) / sum(e.liters for e in priced)
    else:
        last = FuelEntry.query.filter(FuelEntry.kind == "lieferung", FuelEntry.price_per_liter.isnot(None),
                                      FuelEntry.date <= end).order_by(FuelEntry.date.desc(), FuelEntry.id.desc()).first()
        price = last.price_per_liter if last else None

    machines = OrderedDict()
    for e in refuels:
        row = machines.setdefault(e.machine_id, {"machine": e.machine, "liters": Decimal(0), "count": 0, "counters": []})
        row["liters"] += e.liters
        row["count"] += 1
        if e.counter is not None:
            row["counters"].append((e.counter, e.liters))
    for row in machines.values():
        counters = sorted(row["counters"])
        row["first_counter"] = counters[0][0] if counters else None
        row["last_counter"] = counters[-1][0] if counters else None
        row["consumption"] = None
        if len(counters) >= 2 and counters[-1][0] > counters[0][0] and row["machine"].counter_unit:
            liters = sum(float(liters) for _counter, liters in counters[1:])
            value = liters / (counters[-1][0] - counters[0][0])
            row["consumption"] = (value * 100, "l/100 km") if row["machine"].counter_unit == "km" else (value, "l/h")
        row["cost"] = (row["liters"] * price).quantize(Decimal("0.01")) if price else None
    rows = sorted(machines.values(), key=lambda r: (-r["liters"], r["machine"].name))

    return {
        "start": start, "end": end, "machine": machine, "entries": entries,
        "opening": fuel.level(start - timedelta(days=1))[0] if machine is None else None,
        "closing": fuel.level(end)[0] if machine is None else None,
        "delivered": sum((e.liters for e in deliveries), Decimal(0)),
        "delivered_cost": sum((e.liters * e.price_per_liter for e in priced), Decimal(0)).quantize(Decimal("0.01"))
        if priced else None,
        "refueled": sum((e.liters for e in refuels), Decimal(0)),
        "price": Decimal(price).quantize(Decimal("0.001")) if price else None,
        "machines": rows,
    }


def period_label(report):
    start, end = report["start"], report["end"]
    if start == date(start.year, 1, 1) and end == date(start.year, 12, 31):
        return f"Jahr {start.year}"
    return f"{start:%d.%m.%Y} – {end:%d.%m.%Y}"


def filename(report, ext):
    start, end = report["start"], report["end"]
    part = str(start.year) if period_label(report).startswith("Jahr") else f"{start:%Y%m%d}-{end:%Y%m%d}"
    if report["machine"] is not None:
        part += "_" + "".join(c if c.isalnum() else "-" for c in report["machine"].code)
    return f"tankbuch_{part}.{ext}"


def _liters(value, digits=1):
    if value is None:
        return "–"
    return f"{value:,.{digits}f} l".replace(",", "X").replace(".", ",").replace("X", ".")


def _money(value):
    return "–" if value is None else f"{value:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _plain(value, digits=1):
    return f"{value:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


# --- Excel ------------------------------------------------------------------------

HEAD_FILL = PatternFill("solid", fgColor="E7EEDB")


def _head(ws, row):
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_xlsx(report, farm):
    wb = Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    ws.append([f"Tankbuch {farm.get('farm_name') or ''}".strip(), period_label(report)])
    ws["A1"].font = ws["B1"].font = Font(bold=True, size=14)
    if report["machine"] is not None:
        ws.append(["Maschine", report["machine"].name])
    ws.append(["Erstellt am", datetime.now().strftime("%d.%m.%Y %H:%M")])
    ws.append([])
    if report["opening"] is not None:
        ws.append(["Bestand am Anfang (l)", float(report["opening"])])
    ws.append(["Geliefert (l)", float(report["delivered"])])
    if report["delivered_cost"] is not None:
        ws.append(["Kosten der Lieferungen (€)", float(report["delivered_cost"])])
    ws.append(["Getankt (l)", float(report["refueled"])])
    if report["closing"] is not None:
        ws.append(["Bestand am Ende (l)", float(report["closing"])])
    if report["price"] is not None:
        ws.append(["Durchschnittspreis (€/l)", float(report["price"])])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22

    ws = wb.create_sheet("Je Maschine")
    ws.append(["Maschine", "Kennzeichen", "Code", "Tankungen", "Liter", "Zähler von", "Zähler bis", "Einheit",
               "Verbrauch", "Verbrauch Einheit", "Kosten geschätzt (€)"])
    _head(ws, 1)
    for r in report["machines"]:
        m = r["machine"]
        ws.append([m.name, m.plate, m.code, r["count"], float(r["liters"]), r["first_counter"], r["last_counter"],
                   m.unit_short or None, round(r["consumption"][0], 2) if r["consumption"] else None,
                   r["consumption"][1] if r["consumption"] else None, float(r["cost"]) if r["cost"] else None])
    if report["machines"]:
        ws.append(["Summe", None, None, sum(r["count"] for r in report["machines"]), float(report["refueled"])])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, width in enumerate([24, 14, 10, 11, 12, 12, 12, 9, 11, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws = wb.create_sheet("Einträge")
    ws.append(["Datum", "Art", "Maschine", "Kennzeichen", "Zählerstand", "Liter", "Preis (€/l)", "Betrag (€)",
               "Lieferant", "Notiz", "Eingetragen von"])
    _head(ws, 1)
    for e in report["entries"]:
        amount = float((e.liters * e.price_per_liter).quantize(Decimal("0.01"))) if e.price_per_liter else None
        ws.append([e.date, KIND_TEXT.get(e.kind, e.kind), e.machine.name if e.machine else None,
                   e.machine.plate if e.machine else None, e.counter, float(e.liters),
                   float(e.price_per_liter) if e.price_per_liter else None, amount,
                   e.supplier.name if e.supplier else None, e.note, e.user.name if e.user else None])
        ws.cell(row=ws.max_row, column=1).number_format = "DD.MM.YYYY"
    for i, width in enumerate([12, 11, 22, 13, 12, 10, 11, 11, 22, 28, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------------

BASE = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=11.5)
SMALL = ParagraphStyle("small", parent=BASE, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#555555"))
FARM = ParagraphStyle("farm", parent=BASE, fontName="Helvetica-Bold", fontSize=13, leading=16)
TITLE = ParagraphStyle("title", parent=BASE, fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=TA_RIGHT)
RIGHT = ParagraphStyle("right", parent=BASE, alignment=TA_RIGHT)
H2 = ParagraphStyle("h2", parent=BASE, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
HEAD = ParagraphStyle("head", parent=BASE, fontName="Helvetica-Bold", fontSize=8, leading=10)
HEAD_R = ParagraphStyle("head_r", parent=HEAD, alignment=TA_RIGHT)
LINE = colors.HexColor("#999999")
ZEBRA = colors.HexColor("#f3f1ea")


def _p(text, style=BASE):
    return Paragraph(text, style)


def _table(rows, widths, right_cols=(), total=False):
    table = Table(rows, colWidths=widths, repeatRows=1)
    style = [
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in range(1, len(rows) - (1 if total else 0)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), ZEBRA))
    if total:
        style.append(("LINEABOVE", (0, -1), (-1, -1), 0.8, LINE))
    table.setStyle(TableStyle(style))
    return table


def build_pdf(report, farm):
    buf = io.BytesIO()
    farm_name = farm.get("farm_name") or ""
    label = period_label(report)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(18 * mm, 10 * mm, f"{farm_name} · Tankbuch {label} · erstellt {datetime.now():%d.%m.%Y}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Seite {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=f"Tankbuch {label}", author=farm_name)
    width = A4[0] - 36 * mm
    story = []

    left = [_p(escape(farm_name), FARM)] if farm_name else []
    address = "<br/>".join(escape(line) for line in (farm.get("farm_address") or "").splitlines() if line.strip())
    if address:
        left.append(_p(address, SMALL))
    right = [_p("Tankbuch Diesel", TITLE), _p(escape(label), RIGHT)]
    if report["machine"] is not None:
        m = report["machine"]
        right.append(_p(escape(m.name + (f" · {m.plate}" if m.plate else "")), RIGHT))
    head = Table([[left, right]], colWidths=[width * 0.55, width * 0.45])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 8 * mm)]

    # Zusammenfassung
    summary = []
    if report["opening"] is not None:
        summary.append(["Bestand am Anfang", _liters(report["opening"])])
    summary.append(["+ Geliefert", _liters(report["delivered"])
                    + (f"  ({_money(report['delivered_cost'])})" if report["delivered_cost"] is not None else "")])
    summary.append(["− Getankt", _liters(report["refueled"])])
    if report["closing"] is not None:
        summary.append(["Bestand am Ende", _liters(report["closing"])])
    if report["price"] is not None:
        summary.append(["Durchschnittspreis", f"{_plain(report['price'], 3)} €/l"])
    table = Table([[_p(a), _p(b, RIGHT)] for a, b in summary], colWidths=[50 * mm, 50 * mm], hAlign="LEFT")
    table.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#dddddd")),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    story += [table]
    if report["opening"] is not None and report["closing"] is not None:
        story.append(_p("Bestände aus Peilungen, Lieferungen und Tankungen gerechnet.", SMALL))

    # Je Maschine
    story.append(_p("Getankt je Maschine", H2))
    rows = [[_p("Maschine", HEAD), _p("Kennzeichen", HEAD), _p("Tankungen", HEAD_R), _p("Zähler von–bis", HEAD_R),
             _p("Liter", HEAD_R), _p("Verbrauch", HEAD_R), _p("Kosten ca.", HEAD_R)]]
    for r in report["machines"]:
        m = r["machine"]
        counters = f"{r['first_counter']}–{r['last_counter']} {m.unit_short}" if r["first_counter"] is not None else "–"
        cons = f"{_plain(r['consumption'][0])} {r['consumption'][1]}" if r["consumption"] else "–"
        rows.append([_p(escape(m.name)), _p(escape(m.plate or "")), _p(str(r["count"]), RIGHT), _p(counters, RIGHT),
                     _p(_liters(r["liters"]), RIGHT), _p(cons, RIGHT), _p(_money(r["cost"]), RIGHT)])
    if report["machines"]:
        rows.append([_p("<b>Summe</b>"), "", _p(f"<b>{sum(r['count'] for r in report['machines'])}</b>", RIGHT), "",
                     _p(f"<b>{_liters(report['refueled'])}</b>", RIGHT), "", ""])
        story.append(_table(rows, [39 * mm, 23 * mm, 20 * mm, 27 * mm, 20 * mm, 23 * mm, 22 * mm], total=True))
    else:
        story.append(_p("Im Zeitraum nichts getankt.", SMALL))
    if report["price"] is not None and report["machines"]:
        story.append(_p("Kosten ca. = Liter × Durchschnittspreis der Lieferungen. Verbrauch aus den Zählerständen "
                        "(erste Tankung im Zeitraum zählt nur als Ausgangsstand).", SMALL))

    # Einträge
    story.append(_p("Alle Einträge", H2))
    rows = [[_p("Datum", HEAD), _p("Art", HEAD), _p("Maschine / Lieferant", HEAD), _p("Zähler", HEAD_R),
             _p("Liter", HEAD_R), _p("€/l", HEAD_R), _p("Von", HEAD)]]
    for e in report["entries"]:
        who = e.machine.name if e.machine else (e.supplier.name if e.supplier else "")
        if e.note:
            who = f"{escape(who)}<br/><font size=7 color='#555555'>{escape(e.note)}</font>"
        else:
            who = escape(who)
        counter = f"{e.counter} {e.machine.unit_short}" if e.counter is not None and e.machine else ""
        sign = {"tanken": "−", "lieferung": "+", "peilung": "="}.get(e.kind, "")
        rows.append([_p(f"{e.date:%d.%m.%Y}"), _p(KIND_TEXT.get(e.kind, e.kind)), _p(who), _p(counter, RIGHT),
                     _p(sign + _liters(e.liters), RIGHT),
                     _p(_plain(e.price_per_liter, 3) if e.price_per_liter else "", RIGHT),
                     _p(escape(e.user.name) if e.user else "", SMALL)])
    if report["entries"]:
        story.append(_table(rows, [20 * mm, 18 * mm, 52 * mm, 22 * mm, 22 * mm, 14 * mm, 26 * mm]))
    else:
        story.append(_p("Keine Einträge im Zeitraum.", SMALL))

    story += [Spacer(1, 14 * mm), Table([["", ""], [_p("Ort, Datum", SMALL), _p("Unterschrift", SMALL)]],
                                        colWidths=[width * 0.45, width * 0.45], hAlign="LEFT",
                                        style=TableStyle([("LINEABOVE", (0, 1), (0, 1), 0.6, LINE),
                                                          ("LINEABOVE", (1, 1), (1, 1), 0.6, LINE),
                                                          ("TOPPADDING", (0, 0), (-1, 0), 14)]))]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
