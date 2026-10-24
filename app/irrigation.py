"""Beregnung: Import des RainDancer-Wasserbuchs (Excel-Export) und Auswertung je Brunnen und Schlag.

Aufbau der Datei (Bericht „Wasserbuch“, Export „Excel 97-2003“ oder .xlsx):
  Kopf: „Wasserbuch“, „Mandant: …“, „Zeitraum: TT.MM.JJJJ - TT.MM.JJJJ“, dann die Spaltenzeile
  (Datum, Schlag, Regnername, Startzeit, Endzeit, Stunden, Kubikmeter).
  Danach je Brunnen eine Gruppenzeile „Name  (Nr)“ mit Summe, je Tag eine Datumszeile mit Summe und darunter die
  Einsätze (Schlag, Regner, Start, Ende, Stunden, m³). Am Ende „Total“.
Der Name des Mandanten wird nicht gespeichert. Die Zeiten in der Datei sind Ortszeit (Europe/Berlin).

Wie die übrigen Logik-Module ändern die Funktionen nur die Session; committen muss der Aufrufer.
"""
import io
import json
import re
import zipfile
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal

from sqlalchemy import func

from .extensions import db
from .i18n import N_, gettext as _
from .models import Irrigation, Plot, Setting, Well
from .storage import StorageError
from .utils import local_tz

SOURCE = "raindancer"
ALIAS_KEY = "raindancer_plots"               # Schlagname in RainDancer → Schlag-ID (einmal von Hand zugeordnet)
WARN_PERCENT = 80                            # ab hier Hinweis zur genehmigten Jahresmenge
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
HEADERS = {"schlag": "plot", "regnername": "sprinkler", "startzeit": "start", "endzeit": "end", "stunden": "hours",
           "kubikmeter": "m3"}
PERIOD_RE = re.compile(r"Zeitraum:\s*(\d{1,2}\.\d{1,2}\.\d{4})\s*-\s*(\d{1,2}\.\d{1,2}\.\d{4})")
WELL_RE = re.compile(r"^(.*?)\s*\(([^()]*)\)\s*$")


NOT_WATERBOOK = N_("Das ist kein RainDancer-Wasserbuch. Bitte in RainDancer den Bericht „Wasserbuch“ als „Excel 97-2003“ exportieren.")
INCOMPLETE = N_("Die Datei ist unvollständig: Die Summe der Einsätze passt nicht zum Total. Bitte als „Excel 97-2003“ exportieren, nicht als CSV.")


class IrrigationImportError(StorageError):
    pass


def _fail():
    raise IrrigationImportError(_(NOT_WATERBOOK))


# --- Datei lesen ------------------------------------------------------------

def read_rows(data):
    """Excel-Datei (.xls oder .xlsx) → Zeilen mit Python-Werten (Text, Zahl, datetime)."""
    if data.startswith(OLE_MAGIC):
        import xlrd
        try:
            book = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
        except xlrd.XLRDError:
            _fail()
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                row.append(value)
            rows.append(row)
        return rows
    if data.startswith(b"PK\x03\x04"):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except (zipfile.BadZipFile, KeyError, ValueError, OSError):
            _fail()
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    _fail()


def _text(value):
    return re.sub(r"\s+", " ", value).strip() if isinstance(value, str) else ""


def _number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return float(value.strip().replace(".", "").replace(",", "."))
        except ValueError:
            return None
    return None


def _when(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                pass
    return None


def parse(rows):
    """Zeilen → {"period": (von, bis) | None, "events": [...], "total_m3": float | None}."""
    header, cols, period = None, {}, None
    for i, row in enumerate(rows[:20]):
        first = _text(row[0] if row else "")
        m = PERIOD_RE.search(first)
        if m:
            period = tuple(datetime.strptime(x, "%d.%m.%Y").date() for x in m.groups())
        labels = [_text(v).casefold() for v in row]
        if labels and labels[0] == "datum" and "schlag" in labels:
            header = i
            cols = {key: labels.index(label) for label, key in HEADERS.items() if label in labels}
            break
    if header is None or not {"plot", "start", "m3"} <= set(cols):
        _fail()

    def cell(row, key):
        index = cols.get(key)
        return row[index] if index is not None and index < len(row) else None

    events, well, total = [], None, None
    for row in rows[header + 1:]:
        if not row:
            continue
        first, plot = _text(row[0]), _text(cell(row, "plot"))
        start = _when(cell(row, "start"))
        if plot and start:
            m3 = _number(cell(row, "m3"))
            if m3 is None:
                continue
            events.append({"well": well, "plot": plot[:100], "sprinkler": _text(cell(row, "sprinkler"))[:100] or None,
                           "start": start, "end": _when(cell(row, "end")), "hours": _number(cell(row, "hours")),
                           "m3": m3})
        elif first and not plot and not isinstance(row[0], datetime):
            if first.casefold() == "total":
                total = _number(cell(row, "m3"))
            else:                                        # Gruppenzeile: Brunnen
                m = WELL_RE.match(first)
                well = (m.group(1).strip(), m.group(2).strip() or None) if m else (first, None)
                well = (well[0][:100], (well[1] or "")[:30] or None)
    if not events:
        raise IrrigationImportError(_("In der Datei stehen keine Beregnungseinsätze."))
    if total is not None and abs(sum(e["m3"] for e in events) - total) > 1:
        raise IrrigationImportError(_(INCOMPLETE))
    return {"period": period, "events": events, "total_m3": total}


def parse_file(data):
    return parse(read_rows(data))


# --- Zuordnung zu Schlägen --------------------------------------------------------

def _key(name):
    return re.sub(r"\s+", " ", name or "").strip().casefold()


def aliases():
    setting = db.session.get(Setting, ALIAS_KEY)
    try:
        return {k: int(v) for k, v in json.loads(setting.value).items()} if setting and setting.value else {}
    except (ValueError, TypeError, AttributeError):
        return {}


def save_aliases(mapping):
    Setting.put(ALIAS_KEY, json.dumps(mapping, ensure_ascii=False, sort_keys=True) if mapping else None)


def match_plots(names):
    """{Schlagname laut Datei: Schlag oder None} – erst gespeicherte Zuordnung, dann gleicher Name."""
    by_name = {}
    for plot in Plot.query.order_by(Plot.active.desc(), Plot.id):
        by_name.setdefault(_key(plot.name), plot)
    alias = aliases()
    result = {}
    for name in names:
        plot = db.session.get(Plot, alias[_key(name)]) if _key(name) in alias else None
        result[name] = plot or by_name.get(_key(name))
    return result


def plan(parsed):
    """Vorschau: je Schlagname Einsätze, m³ und Zuordnung; je Brunnen m³ und ob er neu ist."""
    plots = {}
    for e in parsed["events"]:
        row = plots.setdefault(e["plot"], {"name": e["plot"], "events": 0, "m3": 0.0})
        row["events"] += 1
        row["m3"] += e["m3"]
    matches = match_plots(plots)
    for name, row in plots.items():
        row["plot"] = matches[name]
    wells = {}
    for e in parsed["events"]:
        name = (e["well"] or ("–", None))[0]
        row = wells.setdefault(name, {"name": name, "number": (e["well"] or (None, None))[1], "m3": 0.0})
        row["m3"] += e["m3"]
    existing = {w.name for w in Well.query}
    for row in wells.values():
        row["new"] = row["name"] not in existing
    days = [e["start"].date() for e in parsed["events"]]
    return {"plots": sorted(plots.values(), key=lambda r: r["name"].casefold()),
            "wells": sorted(wells.values(), key=lambda r: r["name"].casefold()),
            "first": min(days), "last": max(days), "m3": sum(e["m3"] for e in parsed["events"])}


def _utc(naive):
    return naive.replace(tzinfo=local_tz()).astimezone(timezone.utc) if naive else None


def apply(parsed, chosen):
    """Übernehmen. `chosen`: {Schlagname: Schlag-ID oder None} aus der Vorschau (wird als Zuordnung gemerkt).
    Ersetzt alle RainDancer-Einsätze im Zeitraum der Datei. Gibt {"events", "removed", "wells_new"} zurück."""
    alias = aliases()
    for name, plot_id in chosen.items():
        if plot_id and db.session.get(Plot, plot_id):
            alias[_key(name)] = plot_id
        else:
            alias.pop(_key(name), None)
    save_aliases(alias)
    matches = match_plots({e["plot"] for e in parsed["events"]})

    first, last = parsed["period"] or (min(e["start"].date() for e in parsed["events"]),
                                       max(e["start"].date() for e in parsed["events"]))
    since, until = _utc(datetime.combine(first, time())), _utc(datetime.combine(last + timedelta(days=1), time()))
    old = Irrigation.query.filter(Irrigation.source == SOURCE, Irrigation.start_at >= since,
                                  Irrigation.start_at < until)
    removed = old.count()
    for row in old:
        db.session.delete(row)

    wells, new_wells = {w.name: w for w in Well.query}, 0
    for e in parsed["events"]:
        well = None
        if e["well"]:
            name, number = e["well"]
            well = wells.get(name)
            if well is None:
                well = wells[name] = Well(name=name, number=number)
                db.session.add(well)
                new_wells += 1
        plot = matches.get(e["plot"])
        db.session.add(Irrigation(well=well, plot_id=plot.id if plot else None, plot_name=e["plot"],
                                  sprinkler=e["sprinkler"], start_at=_utc(e["start"]), end_at=_utc(e["end"]),
                                  hours=Decimal(str(round(e["hours"], 1))) if e["hours"] is not None else None,
                                  m3=Decimal(str(round(e["m3"], 2))), source=SOURCE))
    db.session.flush()
    return {"events": len(parsed["events"]), "removed": removed, "wells_new": new_wells}


# --- Auswertung -------------------------------------------------------------

def _year_range(year):
    return _utc(datetime(year, 1, 1)), _utc(datetime(year + 1, 1, 1))


def years():
    local = func.extract("year", Irrigation.start_at)
    return sorted({int(y) for (y,) in db.session.query(local).distinct() if y}, reverse=True)


def events(year, plot=None):
    since, until = _year_range(year)
    query = Irrigation.query.filter(Irrigation.start_at >= since, Irrigation.start_at < until)
    if plot is not None:
        query = query.filter(Irrigation.plot_id == plot.id)
    return query.order_by(Irrigation.start_at.desc()).all()


def overview(year):
    """{"wells": [...], "plots": [...], "m3", "hours", "count"} für ein Kalenderjahr."""
    rows = events(year)
    wells, plots = {}, {}
    for e in rows:
        w = wells.setdefault(e.well_id, {"well": e.well, "m3": 0.0, "hours": 0.0, "count": 0})
        w["m3"] += float(e.m3)
        w["hours"] += float(e.hours or 0)
        w["count"] += 1
        key = e.plot_id or f"?{e.plot_name}"
        p = plots.setdefault(key, {"plot": e.plot, "name": e.plot.name if e.plot else e.plot_name, "m3": 0.0,
                                   "hours": 0.0, "count": 0, "last": None})
        p["m3"] += float(e.m3)
        p["hours"] += float(e.hours or 0)
        p["count"] += 1
        p["last"] = max(p["last"], e.start_at) if p["last"] else e.start_at
    for w in wells.values():
        permit = w["well"].permit_m3 if w["well"] else None
        w["percent"] = round(100 * w["m3"] / permit) if permit else None
        w["level"] = (None if w["percent"] is None else "danger" if w["percent"] >= 100
                      else "warning" if w["percent"] >= WARN_PERCENT else "ok")
    for p in plots.values():
        area = float(p["plot"].area_ha) if p["plot"] and p["plot"].area_ha else None
        p["mm"] = p["m3"] / (area * 10) if area else None
    # Brunnen ohne Einsatz im Jahr trotzdem zeigen (Genehmigung eintragen)
    for well in Well.query.order_by(Well.name):
        if well.id not in wells:
            wells[well.id] = {"well": well, "m3": 0.0, "hours": 0.0, "count": 0,
                              "percent": 0 if well.permit_m3 else None, "level": "ok" if well.permit_m3 else None}
    return {"wells": sorted(wells.values(), key=lambda w: (w["well"] is None, w["well"].name if w["well"] else "")),
            "plots": sorted(plots.values(), key=lambda p: -p["m3"]),
            "m3": sum(float(e.m3) for e in rows), "hours": sum(float(e.hours or 0) for e in rows), "count": len(rows),
            "unassigned": sorted({e.plot_name for e in rows if e.plot_id is None})}


def plot_summary(plot, year):
    rows = events(year, plot)
    if not rows:
        return None
    m3 = sum(float(e.m3) for e in rows)
    area = float(plot.area_ha) if plot.area_ha else None
    return {"events": rows, "m3": m3, "hours": sum(float(e.hours or 0) for e in rows),
            "mm": m3 / (area * 10) if area else None, "count": len(rows)}


def mm_by_plot_year(plot_ids, years_):
    """{(plot_id, Jahr): mm} für den Ertragsvergleich."""
    result = {}
    if not plot_ids or not years_:
        return result
    since, until = _year_range(min(years_))[0], _year_range(max(years_))[1]
    for e in Irrigation.query.filter(Irrigation.plot_id.in_(plot_ids), Irrigation.start_at >= since,
                                     Irrigation.start_at < until):
        key = (e.plot_id, e.start_at.astimezone(local_tz()).year)
        result[key] = result.get(key, 0.0) + float(e.m3)
    areas = {p.id: float(p.area_ha) for p in Plot.query.filter(Plot.id.in_(plot_ids)) if p.area_ha}
    return {k: m3 / (areas[k[0]] * 10) for k, m3 in result.items() if k[0] in areas}


def alerts(year=None):
    """Brunnen mit Genehmigung, die WARN_PERCENT erreicht haben (für „Jetzt wichtig“)."""
    year = year or date.today().year
    return [w for w in overview(year)["wells"] if w["level"] in ("warning", "danger") and w["count"]]


def save_permit(well, permit_m3, note=None):
    if permit_m3 is not None and not 0 < permit_m3 <= 100_000_000:
        raise StorageError(_("Bitte eine gültige Menge in m³ angeben."))
    well.permit_m3 = permit_m3
    well.note = (note or "").strip()[:300] or None
