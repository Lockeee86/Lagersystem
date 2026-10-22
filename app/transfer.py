"""Import aus einer Excel-Vorlage (für Zettel-Wirtschaft) und Export je Kartoffeljahr.

Import:
- Jede Zeile beschreibt eine Kiste. Leere Zellen lassen den bisherigen Wert unverändert.
- Die Zeilen werden in der Reihenfolge der Datei ausgeführt (wichtig fürs Stapeln: wer zuerst
  kommt, steht unten).
- `run_import` führt alles aus, entscheidet aber nicht über das Speichern: Der Aufrufer
  committet nur, wenn keine Zeile einen Fehler hatte (alles oder nichts), und macht für die
  Vorschau einfach ein Rollback.

Export: eine Excel-Datei mit Auswertung, Kisten und Bewegungen. Das Kartoffeljahr ist das
Erntejahr; Bewegungen werden über ihre Momentaufnahme (Movement.harvest_date) zugeordnet.
"""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import extract

from . import storage
from .extensions import db
from .i18n import CATALOGS, N_, gettext as _
from .models import FILL_LEVELS, SORT_USAGES, Box, BoxType, Lot, Movement, Plot, SeedClass, Variety, Warehouse
from .storage import StorageError
from .utils import to_local

MAX_ROWS = 5000
TEMPLATE_ROWS = 1000

# Spalten der Vorlage: (Schlüssel, Überschrift, zusätzliche Schreibweisen beim Einlesen)
COLUMNS = [
    ("code", N_("Kiste"), ["kistennummer", "kistennr", "nr", "nummer", "code"]),
    ("variety", N_("Sorte"), []),
    ("plot", N_("Schlag"), ["feld", "acker"]),
    ("harvest_date", N_("Erntedatum"), ["ernte", "datum"]),
    ("fill", N_("Füllstand"), ["fuellstand", "menge"]),
    ("grade", N_("Sortierung"), ["groesse", "größe", "kaliber"]),
    ("warehouse", N_("Lager"), ["halle"]),
    ("row", N_("Reihe"), []),
    ("slot", N_("Platz"), ["stellplatz"]),
    ("remove", N_("Auslagern"), ["raus", "entnommen"]),
    ("note", N_("Notiz"), ["bemerkung", "kommentar"]),
]

# Tafel-Format: eine Zeile pro Partie (Block an der Tafel)
LOT_COLUMNS = [
    ("row", N_("Reihe"), []),
    ("quantity", N_("Menge"), ["anzahl", "kisten", "stueck", "stück"]),
    ("variety", N_("Sorte"), []),
    ("usage", N_("Verwendung"), ["saatspeise", "art"]),
    ("seed_class", N_("Stufe"), ["kategorie", "klasse"]),
    ("grade", N_("Kaliber"), ["sortierung", "groesse", "größe"]),
    ("harvest_date", N_("Datum"), ["erntedatum", "ernte"]),
    ("plot", N_("Schlag"), ["feld", "acker"]),
    ("warehouse", N_("Lager"), ["halle"]),
    ("note", N_("Notiz"), ["bemerkung", "kommentar"]),
]
USAGE_WORDS = {
    "saat": "pflanzgut", "pflanzgut": "pflanzgut", "pfl": "pflanzgut", "sadzeniaki": "pflanzgut",
    "sp": "speise", "speise": "speise", "jadalne": "speise", "futter": "futter", "paszowe": "futter",
    "sonstiges": "sonstiges", "inne": "sonstiges",
}
NEW_COLORS = ["#e6b422", "#c8553d", "#8e5572", "#3d7ea6", "#2a9d8f", "#6a994e", "#bc6c25", "#7f5539",
              "#9d4edd", "#ef476f"]

FILL_WORDS = {
    "voll": 100, "1/1": 100, "¾": 75, "3/4": 75, "dreiviertel": 75, "½": 50, "1/2": 50, "halb": 50,
    "¼": 25, "1/4": 25, "viertel": 25, "leer": 0,
    "pełna": 100, "pelna": 100, "pełne": 100, "pół": 50, "pol": 50, "pusta": 0, "puste": 0,
}
YES_WORDS = {"x", "ja", "j", "yes", "y", "1", "true", "wahr", "tak", "t", "✓", "✔"}


class SemicolonDialect(csv.excel):
    delimiter = ";"


class TransferError(Exception):
    """Fehler, der die ganze Datei betrifft (z. B. falsches Format)."""


class RowError(Exception):
    pass


def _norm(text):
    return re.sub(r"[^0-9a-zäöüßąćęłńóśźż]", "", str(text).casefold())


def _header_aliases(columns):
    """Überschrift (normalisiert) → Spaltenschlüssel, in allen Sprachen."""
    aliases = {}
    for key, label, extra in columns:
        for text in [label, *extra, *(c.get(label, label) for c in CATALOGS.values())]:
            aliases[_norm(text)] = key
    return aliases


# --- Vorlage ----------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F6F2F")
TITLE_FONT = Font(bold=True, size=14)


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template():
    warehouses = Warehouse.query.order_by(Warehouse.name).all()
    varieties = Variety.query.filter_by(active=True).order_by(Variety.name).all()
    plots = Plot.query.filter_by(active=True).order_by(Plot.name).all()
    first_wh = warehouses[0] if warehouses else None

    wb = Workbook()
    ws = wb.active
    ws.title = _("Kisten")
    ws.append([_(label) for _key, label, _extra in COLUMNS])
    _style_header(ws)
    ws.freeze_panes = "A2"
    widths = {"code": 12, "variety": 20, "plot": 20, "harvest_date": 13, "fill": 11, "grade": 14,
              "warehouse": 16, "row": 8, "slot": 8, "remove": 11, "note": 36}
    col = {key: get_column_letter(i + 1) for i, (key, _l, _e) in enumerate(COLUMNS)}
    for key, width in widths.items():
        ws.column_dimensions[col[key]].width = width
    last = TEMPLATE_ROWS + 1
    for r in range(2, last + 1):
        ws[f"{col['code']}{r}"].number_format = "@"
        ws[f"{col['harvest_date']}{r}"].number_format = "DD.MM.YYYY"

    # Auswahllisten auf einem eigenen Blatt
    lists = wb.create_sheet(_("Listen"))
    columns = [
        (_("Sorten"), [v.name for v in varieties]),
        (_("Schläge"), [p.name for p in plots]),
        (_("Füllstand"), [_(label) for _v, label in FILL_LEVELS]),
        (_("Reihen"), [label for _r, label in first_wh.row_labels()] if first_wh else []),
        (_("Lager"), [w.name for w in warehouses]),
        (_("Auslagern"), ["x"]),
    ]
    for c, (title, values) in enumerate(columns, start=1):
        lists.cell(row=1, column=c, value=title)
        lists.column_dimensions[get_column_letter(c)].width = 22
        for r, value in enumerate(values, start=2):
            lists.cell(row=r, column=c, value=value)
    _style_header(lists)

    def validate(key, list_col, values, strict=True):
        if not values:
            return
        letter = get_column_letter(list_col)
        quoted = lists.title.replace("'", "''")
        dv = DataValidation(type="list", formula1=f"='{quoted}'!${letter}$2:${letter}${len(values) + 1}",
                            allow_blank=True, showErrorMessage=strict)
        dv.error = _("Bitte einen Wert aus der Liste wählen.")
        ws.add_data_validation(dv)
        dv.add(f"{col[key]}2:{col[key]}{last}")

    validate("variety", 1, columns[0][1])
    validate("plot", 2, columns[1][1])
    validate("fill", 3, columns[2][1], strict=False)
    validate("row", 4, columns[3][1], strict=False)
    validate("warehouse", 5, columns[4][1])
    validate("remove", 6, ["x"], strict=False)

    # Anleitung
    guide = wb.create_sheet(_("Anleitung"))
    guide.column_dimensions["A"].width = 110
    lines = [
        (_("So funktioniert der Import"), TITLE_FONT),
        (_("Eine Zeile pro Kiste. Nur die Spalte „Kiste“ ist Pflicht (z. B. K-0012 oder einfach 12)."), None),
        (_("Leere Zellen ändern nichts – es wird nur übernommen, was eingetragen ist."), None),
        (_("Sorte, Schlag und Lager müssen so heißen wie im Lagersystem (Auswahlliste nutzen)."), None),
        (_("Füllstand: voll, ¾, ½, ¼ oder leer (auch 100, 75, 50, 25, 0). Bekommt eine leere Kiste eine Sorte ohne Füllstand, gilt sie als voll."), None),
        (_("Reihe + Platz: Die Kiste wird oben auf diesen Stapel gesetzt (bzw. dorthin umgelagert)."), None),
        (_("Die Zeilen werden von oben nach unten ausgeführt: Wer zuerst kommt, steht im Stapel unten."), None),
        (_("Auslagern: ein „x“ nimmt die Kiste aus dem Lager (mit Füllstand „leer“ gilt sie als leer)."), None),
        (_("Lager leer lassen, wenn es nur ein Lager gibt."), None),
        (_("Vor dem Übernehmen zeigt das Lagersystem eine Vorschau. Gibt es Fehler, wird nichts gespeichert."), None),
        ("", None),
        (_("Beispiel"), Font(bold=True)),
        (_("Kiste 12, Sorte Belana, Schlag Acker Nord, Ernte 05.10.2026, voll, Reihe C, Platz 7"), None),
    ]
    for text, font in lines:
        guide.append([text])
        if font:
            guide.cell(row=guide.max_row, column=1).font = font
    wb.move_sheet(guide, offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    return _save(wb)


# --- Einlesen ---------------------------------------------------------------

def _cell(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _find_header(table, columns, required):
    aliases = _header_aliases(columns)
    for idx, row in enumerate(table[:15]):
        found = {i: aliases[_norm(v)] for i, v in enumerate(row) if v is not None and _norm(v) in aliases}
        if required <= set(found.values()) and len(found) >= 2:
            return found, idx
    return None, None


def _rows_from_table(table):
    """Sucht die Kopfzeile und liefert (art, [(Zeilennummer, {schlüssel: wert})]).
    art = "boxes" (Spalte „Kiste“) oder "lots" (Tafel-Format mit „Reihe“ und „Menge“)."""
    kind = "boxes"
    mapping, header_line = _find_header(table, COLUMNS, {"code"})
    if mapping is None:
        kind = "lots"
        mapping, header_line = _find_header(table, LOT_COLUMNS, {"row", "quantity"})
    if mapping is None:
        raise TransferError(_("Keine passende Kopfzeile gefunden (Spalte „Kiste“ oder „Reihe“ + „Menge“). Bitte eine Vorlage verwenden."))
    rows = []
    for idx, row in enumerate(table[header_line + 1:], start=header_line + 2):
        values = {key: _cell(row[i]) for i, key in mapping.items() if i < len(row)}
        if all(v is None for v in values.values()):
            continue
        rows.append((idx, values))
        if len(rows) > MAX_ROWS:
            raise TransferError(_("Zu viele Zeilen (max. {n}).", n=MAX_ROWS))
    return kind, rows


def read_file(filename, data):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 – kaputte/fremde Datei
            raise TransferError(_("Die Excel-Datei konnte nicht gelesen werden.")) from e
        preferred = {_norm(n) for name in ("Kisten", "Partien")
                     for n in [name, *(c.get(name, name) for c in CATALOGS.values())]}
        sheet = next((wb[n] for n in wb.sheetnames if _norm(n) in preferred), wb.worksheets[0])
        table = [list(r) for r in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 20)]
        wb.close()
        return _rows_from_table(table)
    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            dialect = SemicolonDialect
        return _rows_from_table([r for r in csv.reader(io.StringIO(text), dialect)])
    raise TransferError(_("Bitte eine Excel-Datei (.xlsx) oder CSV-Datei hochladen."))


# --- Werte prüfen -----------------------------------------------------------

def _parse_fill(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        number = value * 100 if 0 < value <= 1 else value   # Excel-Prozentzelle: 0,75 = 75 %
        if int(number) in (0, 25, 50, 75, 100) and number == int(number):
            return int(number)
    else:
        text = str(value).casefold().replace("%", "").replace(" ", "")
        if text in FILL_WORDS:
            return FILL_WORDS[text]
        for _v, label in FILL_LEVELS:
            if text == _(label).casefold():
                return _v
        if text.isdigit() and int(text) in (0, 25, 50, 75, 100):
            return int(text)
    raise RowError(_("Füllstand „{value}“ verstehe ich nicht (voll, ¾, ½, ¼, leer).", value=value))


def _season_date(day, month, today=None):
    """Datum ohne Jahr (wie an der Tafel: „23.9“) → laufende Saison.
    Liegt es mehr als zwei Monate in der Zukunft, ist das Vorjahr gemeint."""
    today = today or date.today()
    candidate = date(today.year, month, day)
    if (candidate - today).days > 60:
        candidate = date(today.year - 1, month, day)
    return candidate


def _parse_date(value):
    if value is None or isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel macht aus „23.9“ eine Zahl; „20.10“ würde zu 20.1 – das ist nicht eindeutig
        text = f"{value}"
        if "." in text and len(text.split(".")[1]) == 1 and text.endswith(".1"):
            raise RowError(_("Datum „{value}“ ist nicht eindeutig. Bitte als 20.10.2026 schreiben.", value=value))
    text = str(value).strip().rstrip(".")
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})", text)
    if m:
        try:
            return _season_date(int(m.group(1)), int(m.group(2)))
        except ValueError:
            pass
    raise RowError(_("Datum „{value}“ verstehe ich nicht (z. B. 05.10.2026).", value=value))


def _parse_int(value, label):
    if value is None:
        return None
    try:
        return int(str(value).strip())
    except ValueError:
        raise RowError(_("{field} „{value}“ ist keine Zahl.", field=_(label), value=value)) from None


def _parse_row_label(value, warehouse):
    if value is None:
        return None
    text = str(value).strip().upper()
    for number, label in warehouse.row_labels():
        if text == label or text == str(number):
            return number
    raise RowError(_("Reihe „{value}“ gibt es in {warehouse} nicht.", value=value, warehouse=warehouse.name))


def _is_yes(value):
    return value is not None and str(value).strip().casefold() in YES_WORDS


# --- Import -----------------------------------------------------------------

def run_import(rows, user, create_missing=False):
    """Führt alle Zeilen aus und liefert je Zeile {line, code, actions, errors}.
    Speichert nichts – Commit/Rollback macht der Aufrufer."""
    warehouses = Warehouse.query.order_by(Warehouse.name).all()
    lookup = {
        "variety": {v.name.casefold(): v for v in Variety.query.all()},
        "plot": {p.name.casefold(): p for p in Plot.query.all()},
        "warehouse": {w.name.casefold(): w for w in warehouses},
    }
    default_type = BoxType.query.order_by(BoxType.id).first()
    results = []
    for line, raw in rows:
        result = {"line": line, "code": str(raw.get("code") or ""), "actions": [], "errors": []}
        results.append(result)
        try:
            with db.session.begin_nested():
                _apply_row(raw, user, result, lookup, warehouses, default_type, create_missing)
        except (RowError, StorageError) as e:
            result["errors"].append(str(e))
    return results


def _lookup(lookup, kind, value, label):
    if value is None:
        return None
    found = lookup[kind].get(str(value).strip().casefold())
    if not found:
        raise RowError(_("{field} „{value}“ gibt es nicht.", field=_(label), value=value))
    return found


def _apply_row(raw, user, result, lookup, warehouses, default_type, create_missing):
    code = storage.normalize_code(str(raw.get("code") or ""))
    if not code:
        raise RowError(_("Ungültige Kistennummer."))
    result["code"] = code

    # Erst alles prüfen, dann ändern
    variety = _lookup(lookup, "variety", raw.get("variety"), "Sorte")
    plot = _lookup(lookup, "plot", raw.get("plot"), "Schlag")
    warehouse = _lookup(lookup, "warehouse", raw.get("warehouse"), "Lager")
    harvest = _parse_date(raw.get("harvest_date"))
    fill = _parse_fill(raw.get("fill"))
    grade = raw.get("grade")
    note = raw.get("note")
    remove = _is_yes(raw.get("remove"))
    target_wh = warehouse or (warehouses[0] if warehouses else None)
    row = _parse_row_label(raw.get("row"), target_wh) if raw.get("row") is not None and target_wh else None
    slot = _parse_int(raw.get("slot"), "Platz")
    if (row is None) != (slot is None):
        raise RowError(_("Für das Einlagern bitte Reihe und Platz angeben."))
    if remove and row is not None:
        raise RowError(_("Auslagern und Reihe/Platz gleichzeitig geht nicht."))

    box = Box.query.filter_by(code=code).first()
    if box is None:
        if not create_missing:
            raise RowError(_("{code} gibt es nicht.", code=code))
        if default_type is None:
            raise RowError(_("Es ist noch kein Kistentyp angelegt."))
        box = storage.create_box(code, default_type.id, user)
        result["actions"].append(_("neu angelegt"))

    content_given = any(v is not None for v in (variety, plot, harvest, grade, note))
    if fill is None and variety is not None and box.fill == 0:
        fill = 100   # leere Kiste bekommt eine Sorte → gilt als voll
    if content_given:
        new_fill = fill if fill is not None else box.fill
        storage.update_contents(
            box, user,
            variety_id=variety.id if variety else box.variety_id,
            harvest_date=harvest or box.harvest_date,
            plot_id=plot.id if plot else box.plot_id,
            grade=str(grade) if grade is not None else box.grade,
            fill=new_fill,
            note=str(note) if note is not None else box.note,
        )
        for label, value in ((N_("Sorte"), variety and variety.name), (N_("Schlag"), plot and plot.name),
                             (N_("Ernte"), harvest and harvest.strftime("%d.%m.%Y")),
                             (N_("Sortierung"), grade), (N_("Notiz"), note)):
            if value is not None:
                result["actions"].append(f"{_(label)}: {value}")
    if fill is not None:
        if not content_given:
            storage.set_fill(box, fill, user)
        result["actions"].append(f"{_('Füllstand')}: {box.fill_label}")

    if remove:
        if not box.is_stored:
            raise RowError(_("{code} ist nicht eingelagert.", code=code))
        storage.remove_box(box, user, empty=(fill == 0))
        result["actions"].append(_("ausgelagert"))
    elif row is not None:
        if target_wh is None:
            raise RowError(_("Es ist noch kein Lager angelegt."))
        if box.is_stored and (box.warehouse_id, box.row, box.slot) == (target_wh.id, row, slot):
            result["actions"].append(_("steht schon auf {location}", location=box.location))
        elif box.is_stored:
            storage.move_box(box, target_wh.id, row, slot, user)
            result["actions"].append(_("umgelagert nach {location}", location=box.location))
        else:
            storage.store_box(box, target_wh.id, row, slot, user)
            result["actions"].append(_("eingelagert auf {location}", location=box.location))

    if not result["actions"]:
        result["actions"].append(_("keine Änderung"))


def run_lot_import(rows, user, create_missing=False):
    """Tafel-Format: jede Zeile legt eine neue Partie an (in der Reihenfolge der Datei links → rechts).
    Mit `create_missing` werden unbekannte Sorten, Stufen und Schläge angelegt. Speichert nichts."""
    warehouses = Warehouse.query.order_by(Warehouse.name).all()
    lookup = {
        "variety": {v.name.casefold(): v for v in Variety.query.all()},
        "plot": {p.name.casefold(): p for p in Plot.query.all()},
        "seed_class": {c.name.casefold(): c for c in SeedClass.query.all()},
        "warehouse": {w.name.casefold(): w for w in warehouses},
    }
    default_type = BoxType.query.order_by(BoxType.id).first()
    results = []
    for line, raw in rows:
        result = {"line": line, "code": "", "actions": [], "errors": []}
        results.append(result)
        try:
            with db.session.begin_nested():
                _apply_lot_row(raw, user, result, lookup, warehouses, default_type, create_missing)
        except (RowError, StorageError) as e:
            result["errors"].append(str(e))
    return results


def _lookup_or_create(lookup, kind, value, label, create_missing, result):
    if value is None:
        return None
    name = " ".join(str(value).split())
    found = lookup[kind].get(name.casefold())
    if found:
        return found
    if not create_missing:
        raise RowError(_("{field} „{value}“ gibt es nicht.", field=_(label), value=name))
    if kind == "variety":
        found = Variety(name=name[:100], color=NEW_COLORS[len(lookup[kind]) % len(NEW_COLORS)])
    elif kind == "plot":
        found = Plot(name=name[:100])
    else:
        found = SeedClass(name=name[:20], position=len(lookup[kind]) + 1)
    db.session.add(found)
    db.session.flush()
    lookup[kind][name.casefold()] = found
    result["actions"].append(_("{field} „{value}“ neu angelegt", field=_(label), value=name))
    return found


def _apply_lot_row(raw, user, result, lookup, warehouses, default_type, create_missing):
    warehouse = _lookup(lookup, "warehouse", raw.get("warehouse"), "Lager") or (warehouses[0] if warehouses else None)
    if warehouse is None:
        raise RowError(_("Es ist noch kein Lager angelegt."))
    if default_type is None:
        raise RowError(_("Es ist noch kein Kistentyp angelegt."))
    if raw.get("row") is None:
        raise RowError(_("Bitte eine Reihe wählen."))
    row = _parse_row_label(raw.get("row"), warehouse)
    quantity = storage.parse_quantity(raw.get("quantity"))
    if quantity is None or quantity <= 0:
        raise RowError(_("Menge „{value}“ verstehe ich nicht (z. B. 40 oder 41,75).", value=raw.get("quantity")))
    if raw.get("variety") is None:
        raise RowError(_("Bitte eine Sorte wählen."))
    usage = None
    if raw.get("usage") is not None:
        usage = USAGE_WORDS.get(_norm(raw["usage"]))
        if usage is None:
            raise RowError(_("Verwendung „{value}“ verstehe ich nicht (Saat oder SP).", value=raw["usage"]))
    harvest = _parse_date(raw.get("harvest_date"))
    variety = _lookup_or_create(lookup, "variety", raw.get("variety"), "Sorte", create_missing, result)
    seed_class = _lookup_or_create(lookup, "seed_class", raw.get("seed_class"), "Stufe", create_missing, result)
    plot = _lookup_or_create(lookup, "plot", raw.get("plot"), "Schlag", create_missing, result)
    lot = storage.create_lot(
        user, warehouse.id, row, variety.id, quantity, default_type.id,
        usage=usage or variety.usage, seed_class_id=seed_class.id if seed_class else None,
        grade=str(raw["grade"]) if raw.get("grade") is not None else None,
        plot_id=plot.id if plot else None, harvest_date=harvest,
        note=str(raw["note"]) if raw.get("note") is not None else None)
    result["code"] = lot.location
    result["actions"].append(_("Partie {qty}× {title} angelegt", qty=lot.quantity_label, title=lot.title))


def build_lot_template():
    """Excel-Vorlage im Tafel-Format: eine Zeile pro Partie."""
    warehouses = Warehouse.query.order_by(Warehouse.name).all()
    first_wh = warehouses[0] if warehouses else None
    wb = Workbook()
    ws = wb.active
    ws.title = _("Partien")
    ws.append([_(label) for _key, label, _extra in LOT_COLUMNS])
    _style_header(ws)
    ws.freeze_panes = "A2"
    col = {key: get_column_letter(i + 1) for i, (key, _l, _e) in enumerate(LOT_COLUMNS)}
    for key, width in {"row": 8, "quantity": 9, "variety": 20, "usage": 12, "seed_class": 8, "grade": 10,
                       "harvest_date": 12, "plot": 20, "warehouse": 16, "note": 36}.items():
        ws.column_dimensions[col[key]].width = width
    last = TEMPLATE_ROWS + 1
    for r in range(2, last + 1):
        ws[f"{col['harvest_date']}{r}"].number_format = "@"   # „23.9“ bleibt Text
        ws[f"{col['quantity']}{r}"].number_format = "0.##"

    lists = wb.create_sheet(_("Listen"))
    columns = [
        (_("Sorten"), [v.name for v in Variety.query.filter_by(active=True).order_by(Variety.name)]),
        (_("Verwendung"), [_("Saat"), _("SP"), _("Futter")]),
        (_("Stufen"), [c.name for c in SeedClass.query.filter_by(active=True).order_by(SeedClass.position, SeedClass.name)]),
        (_("Schläge"), [p.name for p in Plot.query.filter_by(active=True).order_by(Plot.name)]),
        (_("Reihen"), [label for _r, label in first_wh.row_labels()] if first_wh else []),
        (_("Lager"), [w.name for w in warehouses]),
    ]
    for c, (title, values) in enumerate(columns, start=1):
        lists.cell(row=1, column=c, value=title)
        lists.column_dimensions[get_column_letter(c)].width = 22
        for r, value in enumerate(values, start=2):
            lists.cell(row=r, column=c, value=value)
    _style_header(lists)
    quoted = lists.title.replace("'", "''")
    for key, list_col, values, strict in (("variety", 1, columns[0][1], False), ("usage", 2, columns[1][1], False),
                                          ("seed_class", 3, columns[2][1], False), ("plot", 4, columns[3][1], False),
                                          ("row", 5, columns[4][1], True), ("warehouse", 6, columns[5][1], True)):
        if not values:
            continue
        letter = get_column_letter(list_col)
        dv = DataValidation(type="list", formula1=f"='{quoted}'!${letter}$2:${letter}${len(values) + 1}",
                            allow_blank=True, showErrorMessage=strict)
        dv.error = _("Bitte einen Wert aus der Liste wählen.")
        ws.add_data_validation(dv)
        dv.add(f"{col[key]}2:{col[key]}{last}")

    guide = wb.create_sheet(_("Anleitung"))
    guide.column_dimensions["A"].width = 110
    for text, font in [
        (_("Tafel übertragen"), TITLE_FONT),
        (_("Eine Zeile pro Block an der Tafel. Pflicht: Reihe, Menge und Sorte."), None),
        (_("Menge in Kisten, Komma für angebrochene Kisten: 41,75 = 41 volle + eine ¾ Kiste."), None),
        (_("Verwendung: Saat oder SP (Speise). Stufe z. B. PB, SE, E, BE, Z. Kaliber z. B. 35/50."), None),
        (_("Datum wie an der Tafel (23.9) oder vollständig (23.09.2026)."), None),
        (_("Die Zeilen einer Reihe von links nach rechts eintragen – so stehen sie dann auch auf der Tafel."), None),
        (_("Vor dem Übernehmen zeigt das Lagersystem eine Vorschau. Gibt es Fehler, wird nichts gespeichert."), None),
    ]:
        guide.append([text])
        if font:
            guide.cell(row=guide.max_row, column=1).font = font
    wb.move_sheet(guide, offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    return _save(wb)


# --- Export -----------------------------------------------------------------

def season_years():
    years = {int(y) for (y,) in db.session.query(extract("year", Movement.harvest_date))
             .filter(Movement.harvest_date.isnot(None)).distinct() if y}
    years |= {int(y) for (y,) in db.session.query(extract("year", Box.harvest_date))
              .filter(Box.harvest_date.isnot(None)).distinct() if y}
    return sorted(years, reverse=True)


def _naive_local(value):
    value = to_local(value)
    return value.replace(tzinfo=None) if value else None


def _sheet(wb, title, headers, rows, widths=None, formats=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    _style_header(ws)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(i) or max(10, min(40, len(str(header)) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width
    for col, fmt in (formats or {}).items():
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = fmt
    return ws


def season_summary(year=None):
    """Auswertung je Erntejahr, Sorte, Verwendung (und Schlag) aus den Momentaufnahmen der Bewegungen.
    Schlüssel: (jahr, sorte, verwendung, schlag) – schlag None = Summe über alle Schläge.
    Geschätzte Ernte: Kisten = höchster Füllstand der Saison × Gewicht des Kistentyps,
    Partien = höchster Bestand der Saison × Gewicht des Kistentyps."""
    query = Movement.query.filter(Movement.harvest_date.isnot(None), Movement.variety_id.isnot(None))
    if year:
        query = query.filter(extract("year", Movement.harvest_date) == year)
    seasons = {}   # ("box"/"lot", id, erntejahr) → Angaben
    for m in query.order_by(Movement.created_at, Movement.id):
        if m.box_id:
            key = ("box", m.box_id, m.harvest_date.year)
            entry = seasons.setdefault(key, {"item": m.box, "year": m.harvest_date.year, "max": 0})
            entry["max"] = max(entry["max"], (m.fill or 0) / 100)
        elif m.lot_id:
            key = ("lot", m.lot_id, m.harvest_date.year)
            entry = seasons.setdefault(key, {"item": m.lot, "year": m.harvest_date.year, "max": 0})
            entry["max"] = max(entry["max"], float(m.quantity_after or 0))
        else:
            continue
        entry["variety"] = m.variety
        entry["usage"] = m.usage or m.variety.usage
        entry["plot"] = m.plot

    groups = {}
    for (kind, _id, _year), entry in seasons.items():
        item = entry["item"]
        full = item.box_type.full_weight_kg if item.box_type else 0
        same_season = (item.harvest_date and item.harvest_date.year == entry["year"]
                       and item.variety_id == entry["variety"].id)
        if kind == "box":
            count = 1
            stock_kg = item.weight_kg if same_season else 0
            stored = 1 if same_season and item.is_stored and not item.is_empty else 0
        else:
            count = entry["max"]
            stock_kg = item.weight_kg if same_season and not item.closed else 0
            stored = float(item.quantity) if same_season and not item.closed else 0
        for plot in (None, entry["plot"] or False):
            g = groups.setdefault((entry["year"], entry["variety"], entry["usage"], plot),
                                  {"boxes": 0, "harvest_kg": 0, "stock_kg": 0, "stored": 0})
            g["boxes"] += count
            g["harvest_kg"] += round(full * entry["max"])
            g["stock_kg"] += stock_kg
            g["stored"] += stored
    return groups


def _num(value):
    """Für Excel: ganze Zahlen als int, sonst auf 2 Stellen gerundet."""
    value = round(float(value), 2)
    return int(value) if value.is_integer() else value


def build_export(year=None):
    wb = Workbook()
    wb.remove(wb.active)
    title = _("Kartoffeljahr {year}", year=year) if year else _("Alle Jahre")

    # Auswertung
    groups = season_summary(year)
    ws = wb.create_sheet(_("Auswertung"))
    ws.append([_("Lagersystem – Auswertung"), title])
    ws["A1"].font = TITLE_FONT
    ws["B1"].font = TITLE_FONT
    ws.append([_("Erstellt am"), _naive_local(datetime.now().astimezone())])
    ws["B2"].number_format = "DD.MM.YYYY HH:MM"
    ws.append([_("Geschätzte Ernte = höchster Füllstand der Kiste in der Saison × Gewicht des Kistentyps.")])
    ws.append([_("Bei Partien: höchster Bestand der Saison (in Kisten) × Gewicht des Kistentyps.")])
    ws.append([])
    headers = [_("Erntejahr"), _("Sorte"), _("Verwendung"), _("Schlag"), _("Kisten"),
               _("Geschätzte Ernte (kg)"), _("Aktueller Bestand (kg)"), _("Kisten im Lager")]

    def table(heading, with_plot):
        ws.append([heading])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(headers)
        _style_header(ws, ws.max_row)
        keys = sorted((k for k in groups if (k[3] is not None) == with_plot),
                      key=lambda k: (-k[0], k[1].name, k[2], k[3].name if k[3] else ""))
        for k in keys:
            g = groups[k]
            ws.append([k[0], k[1].name, _(SORT_USAGES.get(k[2], k[2])),
                       (k[3].name if k[3] else _("ohne Schlag")) if with_plot else None,
                       _num(g["boxes"]), g["harvest_kg"], g["stock_kg"], _num(g["stored"])])
        if not keys:
            ws.append([_("Keine Daten für diesen Zeitraum.")])
        ws.append([])

    table(_("Je Sorte und Verwendung"), with_plot=False)
    table(_("Je Sorte, Verwendung und Schlag"), with_plot=True)
    for i, width in enumerate([12, 22, 14, 22, 9, 22, 24, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Partien (aktueller Stand)
    lots = Lot.query
    if year:
        lots = lots.filter(extract("year", Lot.harvest_date) == year)
    lot_rows = [[
        lot.warehouse.name, lot.warehouse.row_label(lot.row), lot.position, _num(lot.quantity), lot.variety.name,
        lot.usage_label, lot.seed_class.name if lot.seed_class else None, lot.grade, lot.harvest_date,
        lot.plot.name if lot.plot else None, lot.weight_kg, _("abgeschlossen") if lot.closed else _("im Lager"),
        lot.note, _naive_local(lot.updated_at),
    ] for lot in lots.order_by(Lot.closed, Lot.warehouse_id, Lot.row, Lot.position)]
    _sheet(wb, _("Partien"),
           [_("Lager"), _("Reihe"), _("Position"), _("Menge (Kisten)"), _("Sorte"), _("Verwendung"), _("Stufe"),
            _("Kaliber"), _("Datum"), _("Schlag"), _("Gewicht geschätzt (kg)"), _("Status"), _("Notiz"),
            _("Zuletzt geändert")],
           lot_rows, widths={1: 16, 5: 18, 10: 18, 13: 30, 14: 17},
           formats={4: "0.##", 9: "DD.MM.YYYY", 14: "DD.MM.YYYY HH:MM"})

    # Kisten (aktueller Stand)
    boxes = Box.query
    if year:
        boxes = boxes.filter(extract("year", Box.harvest_date) == year)
    box_rows = [[
        b.code, b.variety.name if b.variety and not b.is_empty else "", b.variety.usage_label if b.variety else "",
        b.plot.name if b.plot else "", b.harvest_date, b.fill, b.weight_kg,
        b.box_type.name if b.box_type else "", b.warehouse.name if b.warehouse else "", b.location or "",
        b.grade or "", b.note or "", _naive_local(b.updated_at),
    ] for b in boxes.order_by(Box.code)]
    _sheet(wb, _("Kisten"),
           [_("Kiste"), _("Sorte"), _("Verwendung"), _("Schlag"), _("Erntedatum"), _("Füllstand (%)"),
            _("Gewicht geschätzt (kg)"), _("Kistentyp"), _("Lager"), _("Platz"), _("Sortierung"), _("Notiz"),
            _("Zuletzt geändert")],
           box_rows, widths={1: 10, 2: 18, 4: 18, 5: 12, 12: 30, 13: 17},
           formats={5: "DD.MM.YYYY", 13: "DD.MM.YYYY HH:MM"})

    # Bewegungen
    moves = Movement.query
    if year:
        moves = moves.filter(extract("year", Movement.harvest_date) == year)
    move_rows = [[
        _naive_local(m.created_at), m.subject, m.action_label, m.from_location or "", m.to_location or "",
        m.fill_before, m.fill_after if m.fill_after is not None else m.fill,
        _num(m.quantity_delta) if m.quantity_delta is not None else None,
        _num(m.quantity_after) if m.quantity_after is not None else None,
        m.variety.name if m.variety else "", _(SORT_USAGES.get(m.usage, m.usage)) if m.usage else "",
        m.seed_class.name if m.seed_class else "", m.plot.name if m.plot else "", m.harvest_date,
        m.user.name if m.user else "", m.detail_text or "", m.note or "",
    ] for m in moves.order_by(Movement.created_at, Movement.id)]
    _sheet(wb, _("Bewegungen"),
           [_("Zeitpunkt"), _("Kiste / Partie"), _("Aktion"), _("Von"), _("Nach"), _("Füllstand vorher (%)"),
            _("Füllstand nachher (%)"), _("Menge ± (Kisten)"), _("Bestand danach (Kisten)"), _("Sorte"),
            _("Verwendung"), _("Stufe"), _("Schlag"), _("Erntedatum"), _("Benutzer"), _("Hinweis"), _("Notiz")],
           move_rows, widths={1: 17, 2: 20, 4: 24, 5: 24, 16: 30, 17: 30},
           formats={1: "DD.MM.YYYY HH:MM", 8: "0.##", 9: "0.##", 14: "DD.MM.YYYY"})
    return _save(wb)


def export_filename(year=None):
    stamp = datetime.now().strftime("%Y-%m-%d")
    return f"lager_{year or 'alle-jahre'}_{stamp}.xlsx"

